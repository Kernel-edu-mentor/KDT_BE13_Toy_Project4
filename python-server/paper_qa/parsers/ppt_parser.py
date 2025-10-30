from pptx import Presentation
from typing import List, Dict
import logging

logger = logging.getLogger(__name__)


class PPTParser:
    def parse(self, file_path: str) -> List[Dict]:
        """PPT 파싱 및 요소별 분리"""
        logger.info(f"Parsing PPT: {file_path}")

        prs = Presentation(file_path)
        content_blocks = []

        for slide_idx, slide in enumerate(prs.slides, 1):
            # 슬라이드 텍스트 추출
            slide_text = []

            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text)

            if slide_text:
                content_blocks.append(
                    {
                        "type": "slide",
                        "content": "\n".join(slide_text),
                        "page": slide_idx,
                    }
                )

        logger.info(f"Parsed {len(content_blocks)} slides from PPT")
        return content_blocks


ppt_parser = PPTParser()
